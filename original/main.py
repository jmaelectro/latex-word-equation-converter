from __future__ import annotations

import io
import logging
import os
import re
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, PlainTextResponse, Response, StreamingResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles

from starlette.exceptions import HTTPException as StarletteHTTPException
from starlette.middleware.gzip import GZipMiddleware
from starlette.requests import Request

import math2docx


# ================================================================
# Config
# ================================================================
APP_TITLE = "Ecuaciones a Word (LaTeX → Word OMML)"
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

logger = logging.getLogger("ecuacionesaword")
if not logger.handlers:
    logging.basicConfig(level=logging.INFO)

MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB
Segment = Tuple[str, str]  # ("text" | "inline" | "display", contenido)

# Puedes desactivar reglas específicas del "ejercicio" si no las quieres:
USE_EXERCISE_TWEAKS = True


def _parse_allowed_origins(value: str) -> List[str]:
    return [o.strip() for o in value.split(",") if o.strip()]


def get_allowed_origins() -> List[str]:
    env_val = os.getenv("ALLOWED_ORIGINS", "").strip()
    if env_val:
        return _parse_allowed_origins(env_val)
    return [
        "https://www.ecuacionesaword.com",
        "http://localhost:8000",
        "http://127.0.0.1:8000",
    ]


app = FastAPI(title=APP_TITLE)

app.add_middleware(
    CORSMiddleware,
    allow_origins=get_allowed_origins(),
    allow_credentials=True,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["*"],
)
app.add_middleware(GZipMiddleware, minimum_size=800)

# Static
_static_dir = os.path.join(BASE_DIR, "static")
if os.path.isdir(_static_dir):
    app.mount("/static", StaticFiles(directory=_static_dir), name="static")


# ================================================================
# Blog slugs (CANONICAL) + Aliases (301)
# ================================================================
# Canonical slugs ES -> filename
BLOG_SLUGS_ES: Dict[str, str] = {
    # 6 originales
    "pasar-ecuaciones-chatgpt-word": "blog.html",
    "convertir-documento-latex-word": "blog2.html",
    "ia-chatgpt-a-word-ejercicios": "blog3.html",
    # IMPORTANTE: el slug que enlazas en home/blog es este (Alt+=)
    "pegar-latex-en-word-alt-eq": "blog4.html",
    "que-es-omml-ecuaciones-word": "blog5.html",
    "markdown-con-latex-a-word-docx": "blog6.html",
    # Nuevos (pack SEO)
    "signos-interrogacion-ecuaciones-chatgpt-word": "blog-signos-interrogacion-ecuaciones-chatgpt-word.html",
    "overleaf-latex-a-word-ecuaciones-editables": "blog-overleaf-latex-a-word-ecuaciones-editables.html",
    "omml-vs-mathtype-vs-latex-word-tfg": "blog-omml-vs-mathtype-vs-latex-word-tfg.html",
    "pandoc-ecuaciones-word-no-editables-soluciones": "blog-pandoc-ecuaciones-word-no-editables-soluciones.html",
}

# Canonical slugs EN -> filename
BLOG_SLUGS_EN: Dict[str, str] = {
    # 6 originales
    "copy-chatgpt-equations-word": "blog-en-1.html",
    "convert-latex-document-to-word": "blog-en-2.html",
    "use-ai-equations-to-word-exercises": "blog-en-3.html",
    "paste-latex-into-word-alt-eq": "blog-en-4.html",
    "what-is-omml-word-equations": "blog-en-5.html",
    "markdown-latex-to-word-docx": "blog-en-6.html",
    # Nuevos (pack SEO)
    "question-marks-chatgpt-equations-word": "blog-en-question-marks-chatgpt-equations-word.html",
    "overleaf-latex-to-word-editable-equations": "blog-en-overleaf-latex-to-word-editable-equations.html",
    "omml-vs-mathtype-vs-latex-word-thesis": "blog-en-omml-vs-mathtype-vs-latex-word-thesis.html",
    "pandoc-math-to-word-omml-troubleshooting": "blog-en-pandoc-math-to-word-omml-troubleshooting.html",
}

# Aliases ES (slugs viejos) -> canonical
BLOG_ALIASES_ES: Dict[str, str] = {
    # Antes usabas este slug para el post de Alt+=
    "pegar-latex-editor-ecuaciones-word": "pegar-latex-en-word-alt-eq",

    # FIX: tu índice enlaza a este slug corto, pero el canonical en main.py es el largo.
    # Con esto, /blog/signos-interrogacion-chatgpt-word redirige (301) al slug canonical que ya tienes.
    "signos-interrogacion-chatgpt-word": "signos-interrogacion-ecuaciones-chatgpt-word",
}


# Aliases EN (slugs viejos) -> canonical
BLOG_ALIASES_EN: Dict[str, str] = {
    # En tu web llegó a existir este:
    "copy-chatgpt-equations-to-word": "copy-chatgpt-equations-word",
    # Estos los tenías en versiones previas del main.py:
    "use-ai-to-solve-exercises-and-export-to-word": "use-ai-equations-to-word-exercises",
    "paste-latex-into-word-equation-editor": "paste-latex-into-word-alt-eq",
}


# ================================================================
# Helpers
# ================================================================
def _read_text_file(path: str, default: Optional[str] = None) -> str:
    try:
        with open(path, "r", encoding="utf-8") as f:
            return f.read()
    except FileNotFoundError:
        if default is not None:
            return default
        raise


def read_html_file(filename: str) -> str:
    path = os.path.join(BASE_DIR, filename)
    return _read_text_file(path)


def _now_lastmod_iso() -> str:
    return datetime.now(timezone.utc).date().isoformat()


def _abs_url(path: str) -> str:
    # Canonical host (ajusta si cambias dominio)
    return f"https://www.ecuacionesaword.com{path}"


def _sitemap_url_entry(loc: str, lastmod: str, changefreq: str, priority: str) -> str:
    return (
        "  <url>\n"
        f"    <loc>{loc}</loc>\n"
        f"    <lastmod>{lastmod}</lastmod>\n"
        f"    <changefreq>{changefreq}</changefreq>\n"
        f"    <priority>{priority}</priority>\n"
        "  </url>\n"
    )


def generate_sitemap_xml() -> str:
    """
    Sitemap XML (válido) generado desde las rutas reales.
    """
    lastmod = _now_lastmod_iso()
    urls: List[str] = []

    # Home
    urls.append(_sitemap_url_entry(_abs_url("/"), lastmod, "weekly", "1.0"))
    urls.append(_sitemap_url_entry(_abs_url("/en"), lastmod, "weekly", "0.8"))

    # Blog index
    urls.append(_sitemap_url_entry(_abs_url("/blog"), lastmod, "weekly", "0.8"))
    urls.append(_sitemap_url_entry(_abs_url("/en/blog"), lastmod, "weekly", "0.7"))

    # Blog posts ES
    for slug in BLOG_SLUGS_ES.keys():
        urls.append(_sitemap_url_entry(_abs_url(f"/blog/{slug}"), lastmod, "monthly", "0.6"))

    # Blog posts EN
    for slug in BLOG_SLUGS_EN.keys():
        urls.append(_sitemap_url_entry(_abs_url(f"/en/blog/{slug}"), lastmod, "monthly", "0.5"))

    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">\n'
        + "".join(urls)
        + "</urlset>\n"
    )


# ================================================================
# 1) Normalización / prettify (tu lógica existente)
# ================================================================
def normalize_math_text(text: str) -> str:
    t = text
    t = re.sub(r"q([1-4])\s*\(", r"q_\1(", t)          # q1( -> q_1(
    t = re.sub(r"\bD([1-4])\b", r"D_\1", t)            # D1 -> D_1
    for var in ("x", "y", "z"):
        for exp in ("2", "3", "4"):
            t = re.sub(rf"{var}{exp}\b", rf"{var}^{exp}", t)  # x2 -> x^2
    return t


def _prettify_paragraphs_for_exercise(paragraph_texts: List[str]) -> List[str]:
    out: List[str] = []
    for text in paragraph_texts:
        s = normalize_math_text(text)
        stripped = s.strip()
        if stripped == "":
            continue

        if stripped.lower() == "actividad2grupal":
            out.append("Actividad 2 (trabajo grupal)")
            continue

        if all(sym in stripped for sym in ["q_1(x,y,z)", "q_2(x,y,z)", "q_3(x,y,z)", "q_4(x,y,z)"]):
            out.append("$$ q_1(x,y,z) = 2x^2 + 2y^2 + 2z^2 + 2xy + 2xz $$")
            out.append("$$ q_2(x,y,z) = x^2 - y^2 + z^2 + 2xy $$")
            out.append("$$ q_3(x,y,z) = 2x^2 - y^2 + 2z^2 + 4xy - 4yz $$")
            out.append("$$ q_4(x,y,z) = 4x^2 + 2y^2 + z^2 + 2yz $$")
            continue

        if stripped.startswith("Interpretamos que") and "garantizar beneficios" in stripped:
            out.append(
                "Interpretamos que “garantizar beneficios” significa que el beneficio "
                "$q(x,y,z)$ sea positivo para todo $(x,y,z)\\neq(0,0,0)$, es decir, "
                "que la forma cuadrática sea definida positiva."
            )
            continue

        if (
            "Escribimos cada forma como q(x)=xT" in stripped
            or "Escribimos cada forma como q(x)=xTA" in stripped
            or "Escribimos cada forma como q(x)=xTA x(\\mathbf x)" in stripped
            or ("Escribimos cada forma como q(x)=xTAx(\\mathbf x)=\\mathbf x^T A\\mathbf x" in stripped)
        ):
            out.append("Escribimos cada forma como $q(\\mathbf x) = \\mathbf x^T A\\, \\mathbf x$, con $\\mathbf x = (x,y,z)^T$.")
            continue

        if "coeficiente del término cruzado" in stripped:
            out.append("Recordando que el coeficiente del término cruzado $2x_i y_j$ se reparte como $a_{ij}=a_{ji}=1$:")
            continue

        if "Forma q1" in stripped or "q1q_1q1" in stripped:
            out.append("Forma $q_1$:")
            continue
        if "Forma q2" in stripped or "q2q_2q2" in stripped:
            out.append("Forma $q_2$:")
            continue
        if "Forma q3" in stripped or "q3q_3q3" in stripped:
            out.append("Forma $q_3$:")
            continue
        if "Forma q4" in stripped or "q4q_4q4" in stripped:
            out.append("Forma $q_4$:")
            continue

        if "Criterio de Sylvester" in stripped:
            out.append("2. Criterio de Sylvester")
            continue
        if "Una matriz simétrica" in stripped and "definida positiva" in stripped:
            out.append("Una matriz simétrica $A$ es definida positiva si todos sus menores principales líderes son positivos:")
            continue

        if "D_1>0" in stripped and "D_2>0" in stripped and "D_3>0" in stripped:
            out.append("\\begin{aligned}\nD_1 &> 0,\\\\\nD_2 &> 0,\\\\\nD_3 &> 0.\n\\end{aligned}")
            continue

        d_terms = list(re.finditer(r"D_[1-4][^D]*", stripped))
        if len(d_terms) >= 2:
            for m in d_terms:
                term = m.group(0).strip().strip(",")
                if term:
                    out.append(f"$$ {term} $$")
            continue

        if "A1A_1A1" in stripped or "A1A1" in stripped:
            out.append("⇒ $A_1$ es definida positiva ⇒ $q_1$ definida positiva.")
            continue
        if "A4A_4A4" in stripped or "A4A4" in stripped:
            out.append("⇒ $A_4$ es definida positiva ⇒ $q_4$ definida positiva.")
            continue

        if "detA2" in stripped or "det A_2" in stripped:
            out.append("$$ \\det A_2 = -2 < 0 $$")
            continue
        if "detA3" in stripped or "det A_3" in stripped:
            out.append("$$ \\det A_3 = -20 < 0 $$")
            continue

        if "q3q_3q3" in stripped and "indefinida" in stripped:
            out.append("⇒ $q_3$ también es indefinida.")
            continue

        out.append(stripped)

    return out


def prettify_paragraphs(paragraph_texts: List[str]) -> List[str]:
    generic: List[str] = []
    for text in paragraph_texts:
        s = normalize_math_text(text)
        if s.strip():
            generic.append(s.strip())
    if not USE_EXERCISE_TWEAKS:
        return generic
    return _prettify_paragraphs_for_exercise(generic)


# ================================================================
# 2) Parsing LaTeX
# ================================================================
def new_paragraph(doc: Document, align: Optional[Any] = None):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    return p


def add_math_safe(paragraph, latex: str) -> None:
    try:
        math2docx.add_math(paragraph, latex)
    except Exception as exc:  # noqa: BLE001
        logger.warning("Fallo convirtiendo LaTeX a OMML: %s", exc)
        paragraph.add_run(latex)


def parse_math_segments(text: str) -> List[Segment]:
    segments: List[Segment] = []
    buf: List[str] = []

    def flush_text() -> None:
        if buf:
            segments.append(("text", "".join(buf)))
            buf.clear()

    i = 0
    n = len(text)
    while i < n:
        if text.startswith("$$", i):
            flush_text()
            end = text.find("$$", i + 2)
            if end == -1:
                buf.append(text[i:])
                break
            latex = text[i + 2 : end]
            segments.append(("display", latex.strip()))
            i = end + 2
            continue

        if text.startswith(r"\[", i):
            flush_text()
            end = text.find(r"\]", i + 2)
            if end == -1:
                buf.append(text[i:])
                break
            latex = text[i + 2 : end]
            segments.append(("display", latex.strip()))
            i = end + 2
            continue

        if text[i] == "$":
            flush_text()
            end = text.find("$", i + 1)
            if end == -1:
                buf.append(text[i:])
                break
            latex = text[i + 1 : end]
            segments.append(("inline", latex.strip()))
            i = end + 1
            continue

        buf.append(text[i])
        i += 1

    flush_text()
    return segments


def split_into_paragraph_descriptors(text: str) -> List[Dict[str, Any]]:
    segments = parse_math_segments(text)
    paragraphs: List[Dict[str, Any]] = []
    current_chunks: List[Segment] = []

    for kind, value in segments:
        if kind in ("text", "inline"):
            if value:
                current_chunks.append((kind, value))
        elif kind == "display":
            if current_chunks:
                paragraphs.append({"type": "inline", "chunks": current_chunks})
                current_chunks = []
            paragraphs.append({"type": "display", "latex": value})

    if current_chunks:
        paragraphs.append({"type": "inline", "chunks": current_chunks})

    return paragraphs


def split_long_latex_equation(latex: str, max_len: int = 120) -> List[str]:
    latex = latex.strip()
    if not latex:
        return []
    if len(latex) <= max_len:
        return [latex]

    if "\\\\" in latex:
        return [p.strip() for p in latex.split("\\\\") if p.strip()]

    if "=" in latex:
        tokens = [t.strip() for t in latex.split("=")]
        eqs: List[str] = []
        current = tokens[0]
        for tok in tokens[1:]:
            candidate = current + " = " + tok
            if len(candidate) > max_len and current:
                eqs.append(current)
                current = tok
            else:
                current = candidate
        if current:
            eqs.append(current)
        return eqs

    if "," in latex:
        tokens = [t.strip() for t in latex.split(",")]
        eqs2: List[str] = []
        current2 = ""
        for tok in tokens:
            candidate = (current2 + ", " if current2 else "") + tok
            if len(candidate) > max_len and current2:
                eqs2.append(current2)
                current2 = tok
            else:
                current2 = candidate
        if current2:
            eqs2.append(current2)
        return eqs2

    return [latex[i : i + max_len].strip() for i in range(0, len(latex), max_len)]


def add_aligned_block(doc: Document, aligned_block: str) -> None:
    content = aligned_block.strip()
    if content.startswith(r"\begin{aligned}"):
        content = content[len(r"\begin{aligned}") :]
    if content.endswith(r"\end{aligned}"):
        content = content[: -len(r"\end{aligned}")]
    content = content.strip()
    if not content:
        return
    rows = [row.strip() for row in content.split(r"\\") if row.strip()]
    for row in rows:
        row_no_amp = row.replace("&", "")
        p = new_paragraph(doc, WD_ALIGN_PARAGRAPH.CENTER)
        add_math_safe(p, row_no_amp)


def build_document_from_paragraphs(paragraph_texts: List[str]) -> Document:
    out_doc = Document()
    if not paragraph_texts:
        new_paragraph(out_doc)
        return out_doc

    in_display_block = False
    display_block_delim: Optional[str] = None  # "$$" o "\["
    display_lines: List[str] = []

    in_aligned_block = False
    aligned_lines: List[str] = []

    matrix_envs = ["vmatrix", "bmatrix", "pmatrix", "matrix"]

    for raw_text in paragraph_texts:
        text = normalize_math_text(raw_text)
        stripped = text.strip()

        # aligned multi-línea
        if in_aligned_block:
            aligned_lines.append(text)
            if r"\end{aligned}" in text:
                add_aligned_block(out_doc, "\n".join(aligned_lines))
                in_aligned_block = False
                aligned_lines = []
            continue

        if r"\begin{aligned}" in stripped:
            if r"\end{aligned}" in stripped:
                add_aligned_block(out_doc, stripped)
            else:
                in_aligned_block = True
                aligned_lines = [stripped]
            continue

        # display multi-línea $$ ... $$  o \[ ... \]
        if in_display_block:
            if (
                (display_block_delim == "$$" and stripped == "$$")
                or (display_block_delim == r"\[" and stripped == r"\]")
            ):
                latex = "\n".join(display_lines).strip()
                for part in split_long_latex_equation(latex):
                    p = new_paragraph(out_doc, WD_ALIGN_PARAGRAPH.CENTER)
                    add_math_safe(p, part)
                in_display_block = False
                display_block_delim = None
                display_lines = []
            else:
                display_lines.append(text)
            continue

        if stripped == "$$":
            in_display_block = True
            display_block_delim = "$$"
            display_lines = []
            continue

        if stripped == r"\[":
            in_display_block = True
            display_block_delim = r"\["
            display_lines = []
            continue

        # matrices en una línea
        handled_matrix = False
        for env in matrix_envs:
            begin = rf"\begin{{{env}}}"
            end = rf"\end{{{env}}}"
            if begin in stripped and end in stripped and "$" not in stripped:
                p = new_paragraph(out_doc, WD_ALIGN_PARAGRAPH.CENTER)
                add_math_safe(p, stripped)
                handled_matrix = True
                break
        if handled_matrix:
            continue

        # normal
        descriptors = split_into_paragraph_descriptors(text)
        if not descriptors:
            new_paragraph(out_doc)
            continue

        for desc in descriptors:
            if desc["type"] == "inline":
                p = new_paragraph(out_doc)
                for kind, value in desc["chunks"]:
                    if kind == "text":
                        p.add_run(value)
                    elif kind == "inline":
                        add_math_safe(p, value)
            elif desc["type"] == "display":
                for part in split_long_latex_equation(desc["latex"].strip()):
                    p = new_paragraph(out_doc, WD_ALIGN_PARAGRAPH.CENTER)
                    add_math_safe(p, part)

    # cierre si termina dentro de display/aligned
    if in_display_block and display_lines:
        latex = "\n".join(display_lines).strip()
        for part in split_long_latex_equation(latex):
            p = new_paragraph(out_doc, WD_ALIGN_PARAGRAPH.CENTER)
            add_math_safe(p, part)

    if in_aligned_block and aligned_lines:
        add_aligned_block(out_doc, "\n".join(aligned_lines))

    return out_doc


# ================================================================
# Routes: páginas
# ================================================================
@app.get("/", response_class=HTMLResponse)
async def home() -> HTMLResponse:
    try:
        return HTMLResponse(read_html_file("index.html"))
    except FileNotFoundError:
        return HTMLResponse("<h1>index.html not found</h1>", status_code=404)


@app.get("/en", response_class=HTMLResponse)
async def home_en() -> HTMLResponse:
    try:
        return HTMLResponse(read_html_file("index-en.html"))
    except FileNotFoundError:
        return HTMLResponse("<h1>index-en.html not found</h1>", status_code=404)


@app.get("/blog", response_class=HTMLResponse)
async def blog_index_es() -> HTMLResponse:
    try:
        return HTMLResponse(read_html_file("blog-index.html"))
    except FileNotFoundError:
        return HTMLResponse("<h1>blog-index.html not found</h1>", status_code=404)


@app.get("/en/blog", response_class=HTMLResponse)
async def blog_index_en() -> HTMLResponse:
    try:
        return HTMLResponse(read_html_file("blog-index-en.html"))
    except FileNotFoundError:
        return HTMLResponse("<h1>blog-index-en.html not found</h1>", status_code=404)


# Redirects legacy numeric
@app.get("/blog2")
async def blog2_redirect() -> RedirectResponse:
    return RedirectResponse(url="/blog/convertir-documento-latex-word", status_code=301)


@app.get("/blog3")
async def blog3_redirect() -> RedirectResponse:
    return RedirectResponse(url="/blog/ia-chatgpt-a-word-ejercicios", status_code=301)


@app.get("/blog4")
async def blog4_redirect() -> RedirectResponse:
    # Antes podía existir /blog4 -> el post Alt+=
    return RedirectResponse(url="/blog/pegar-latex-en-word-alt-eq", status_code=301)


@app.get("/blog5")
async def blog5_redirect() -> RedirectResponse:
    return RedirectResponse(url="/blog/que-es-omml-ecuaciones-word", status_code=301)


@app.get("/blog6")
async def blog6_redirect() -> RedirectResponse:
    return RedirectResponse(url="/blog/markdown-con-latex-a-word-docx", status_code=301)


@app.get("/blog/{slug}", response_class=HTMLResponse)
async def blog_post_es(slug: str) -> HTMLResponse:
    # Alias -> canonical
    if slug in BLOG_ALIASES_ES:
        return RedirectResponse(url=f"/blog/{BLOG_ALIASES_ES[slug]}", status_code=301)

    filename = BLOG_SLUGS_ES.get(slug)
    if not filename:
        raise HTTPException(status_code=404, detail="Blog post not found")

    try:
        return HTMLResponse(read_html_file(filename))
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Blog post file not found")


@app.get("/en/blog/{slug}", response_class=HTMLResponse)
async def blog_post_en(slug: str) -> HTMLResponse:
    if slug in BLOG_ALIASES_EN:
        return RedirectResponse(url=f"/en/blog/{BLOG_ALIASES_EN[slug]}", status_code=301)

    filename = BLOG_SLUGS_EN.get(slug)
    if not filename:
        raise HTTPException(status_code=404, detail="Blog post not found")

    try:
        return HTMLResponse(read_html_file(filename))
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Blog post file not found")


@app.get("/robots.txt")
async def robots_txt() -> Response:
    default = "User-agent: *\nAllow: /\nSitemap: https://www.ecuacionesaword.com/sitemap.xml\n"
    path = os.path.join(BASE_DIR, "robots.txt")
    content = _read_text_file(path, default=default)
    return Response(content=content, media_type="text/plain")


@app.get("/sitemap.xml")
async def sitemap_xml() -> Response:
    # Generado siempre en XML válido (no dependes del fichero)
    content = generate_sitemap_xml()
    return Response(content=content, media_type="application/xml")


@app.get("/healthz")
async def healthz() -> PlainTextResponse:
    return PlainTextResponse("ok")


# ================================================================
# Convert endpoint
# ================================================================
def _extract_text_lines_from_docx(file_bytes: bytes) -> List[str]:
    doc = Document(io.BytesIO(file_bytes))
    lines: List[str] = []
    for p in doc.paragraphs:
        # Mantener líneas aunque sean vacías para separar bloques
        lines.append(p.text if p.text is not None else "")
    return lines


def _extract_text_lines_from_txt(file_bytes: bytes) -> List[str]:
    text = file_bytes.decode("utf-8", errors="replace")
    return text.splitlines()


@app.post("/convert")
async def convert(file: UploadFile = File(...)) -> StreamingResponse:
    filename = (file.filename or "").lower().strip()
    if not filename:
        raise HTTPException(status_code=400, detail="Missing filename")

    content = await file.read()
    if len(content) > MAX_FILE_SIZE_BYTES:
        raise HTTPException(status_code=413, detail="File too large (max 5MB)")

    if filename.endswith(".docx"):
        paragraph_texts = _extract_text_lines_from_docx(content)
    elif filename.endswith(".txt"):
        paragraph_texts = _extract_text_lines_from_txt(content)
    else:
        raise HTTPException(status_code=400, detail="Unsupported file type. Use .docx or .txt")

    cleaned = prettify_paragraphs(paragraph_texts)
    out_doc = build_document_from_paragraphs(cleaned)

    out = io.BytesIO()
    out_doc.save(out)
    out.seek(0)

    download_name = "ecuaciones-a-word.docx"
    headers = {
        "Content-Disposition": f'attachment; filename="{download_name}"',
        "Cache-Control": "no-store",
    }

    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=headers,
    )


# ================================================================
# 404 handler (HTML en vez de JSON)
# ================================================================
@app.exception_handler(StarletteHTTPException)
async def custom_http_exception_handler(request: Request, exc: StarletteHTTPException):
    if exc.status_code == 404:
        is_en = str(request.url.path).startswith("/en")
        if is_en:
            html = f"""
            <h1>404 · Page not found</h1>
            <p>The URL <code>{request.url.path}</code> does not exist.</p>
            <p><a href="/en">Go to the converter</a> · <a href="/en/blog">Read the guides</a></p>
            """
        else:
            html = f"""
            <h1>404 · Página no encontrada</h1>
            <p>No existe la URL <code>{request.url.path}</code>.</p>
            <p><a href="/">Ir al conversor</a> · <a href="/blog">Ver las guías</a></p>
            """
        return HTMLResponse(html, status_code=404)

    return PlainTextResponse(str(exc.detail), status_code=exc.status_code)

